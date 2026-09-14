# -*- coding: utf-8 -*-
"""Файл сбора исполнителей (docs/suppliers-template.xlsx) → строки для импорта.

    python ops/suppliers-xlsx-to-json.py заполненный.xlsx > rows.json

Дальше — сухой прогон в API (по умолчанию dryRun: ничего не пишет, только
отчёт: сколько создастся, сколько обновится, какие строки отвергнуты):

    curl -s -X POST https://kerektap.kz/api/admin/suppliers/import \
      -H "Authorization: Bearer <токен админа>" -H "Content-Type: application/json" \
      --data-binary @rows.json

Когда отчёт устраивает — добавить в rows.json "dryRun": false и повторить.

Скрипт переводит название категории в код (slug) по листу «Справочники»
того же файла, чтобы человеку не приходилось знать коды. Строки без телефона
пропускаются; строки, где проверка не «ок», выводятся в stderr и в файл не
попадают — их чинят в таблице, а не в JSON.
"""
import json
import sys

from openpyxl import load_workbook

if len(sys.argv) < 2:
    sys.exit("использование: python ops/suppliers-xlsx-to-json.py файл.xlsx > rows.json")

wb = load_workbook(sys.argv[1], data_only=True)
ref = wb["Справочники"]
slug_by_name = {}
for r in ref.iter_rows(min_row=2, min_col=2, max_col=3, values_only=True):
    name, slug = r
    if name and slug:
        slug_by_name[str(name).strip().lower()] = str(slug).strip()

ws = wb["Исполнители"]
rows, skipped = [], []
for i, r in enumerate(ws.iter_rows(min_row=2, max_col=9, values_only=True), 2):
    _, phone, company, city, category, _src, _link, _note, check = r
    if phone in (None, ""):
        continue
    if str(check).strip() != "ок":
        skipped.append((i, phone, check))
        continue
    slug = slug_by_name.get(str(category).strip().lower())
    if not slug:
        skipped.append((i, phone, f"категория не из справочника: {category}"))
        continue
    row = {"phone": str(phone).strip(), "city": str(city).strip(), "categorySlug": slug}
    if company:
        row["companyName"] = str(company).strip()
    rows.append(row)

for i, phone, why in skipped:
    print(f"пропущена строка {i} ({phone}): {why}", file=sys.stderr)
print(f"строк к импорту: {len(rows)}, пропущено: {len(skipped)}", file=sys.stderr)

json.dump({"rows": rows, "dryRun": True}, sys.stdout, ensure_ascii=False, indent=1)
